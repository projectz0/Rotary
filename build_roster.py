#!/usr/bin/env python3
"""
Gera Rotary_Frequencia_Mestre.xlsx — Rotary Club de Feira de Santana
Ano Rotário 2026/27. Arquivo único e cumulativo (não criar um novo por mês).

Cada linha de `rows`: (Nº, Nome, Dispensado, d1..d8, Observações)
d1-d4 = julho (07/07, 14/07, 21/07, 28/07) — já fechado e confirmado.
d5-d8 = agosto (04/08, 11/08, 18/08, 25/08) — fechado e confirmado com o
        Jefferson em 18/09/2026, lido a partir da lista "Checada" (marca-texto).
Códigos: P (presente), AJ (falta justificada, soma ponto), A (falta sem
justificativa, 0 ponto), "-" (não constava na lista naquela data).
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATE_LABELS = [
    "07/07/2026", "14/07/2026", "21/07/2026", "28/07/2026",
    "04/08/2026", "11/08/2026", "18/08/2026", "25/08/2026",
]
N_DATES = len(DATE_LABELS)

# (Nº, Nome, Dispensado, d1, d2, d3, d4, d5, d6, d7, d8, Observações)
rows = [
    (1, "Adauto Alves Franco Junior", None, "A", "P", "P", "P", "P", "P", "P", "P", None),
    (2, "André Campana Correia Leite", None, "A", "P", "P", "P", "P", "P", "P", "P", None),
    (3, "Albertony Santos Assis", None, "P", "P", "A", "P", "P", "A", "P", "P", "Nome grafado 'SSSIS - Tony' na 1a lista"),
    (4, "Alpiniano Reis Oliveira Filho", None, "A", "P", "P", "P", "P", "P", "P", "P", None),
    (5, "André Fernandes Dórea", "*", "A", "P", "P", "P", "P", "P", "P", "P", None),
    (6, "Antônio Washington F. de Almeida", "*", "P", "P", "P", "P", "A", "A", "A", "P", None),
    (7, "Antônio José de Carvalho Monteiro", None, "P", "P", "A", "P", "P", "P", "P", "P", "Nome com apelido '-Toni' na 1a lista"),
    (8, "Alexandre Tarquínio Lara Medrado", None, "P", "A", "A", "P", "A", "A", "A", "A", None),
    (9, "Bruno de Nunes Silva", None, "P", "P", "P", "P", "P", "P", "P", "P", None),
    (10, "Carlos Olimpio Ferraz Ribeiro Junior", None, "A", "A", "A", "A", "A", "A", "A", "A", None),
    (11, "Claudiana de Souza Santos Carvalho", None, "A", "P", "A", "A", "A", "A", "P", "A", None),
    (12, "Cleber Chagas Gonçalves", None, "A", "P", "A", "A", "A", "A", "A", "P", None),
    (13, "Cristiano Almeida de Oliveira", None, "P", "P", "A", "A", "P", "A", "A", "A", None),
    (14, "Cristina Fernandes da Cruz Dórea", None, "P", "A", "P", "P", "A", "A", "A", "A", None),
    (15, "Dázio Brasileiro Filho", "*", "P", "P", "P", "P", "P", "P", "P", "P", None),
    (16, "Deonilso Aparecido Bueno de Oliveira", "*", "A", "A", "A", "P", "A", "P", "P", "P", None),
    (17, "Dion Luciano Vital", None, "P", "A", "P", "P", "P", "P", "P", "P", None),
    (18, "Edivaldo Rodrigues Santana", None, "P", "P", "P", "P", "P", "P", "P", "P", None),
    (19, "Eliana Mattos de Amorim Bueno", None, "P", "AJ", "AJ", "P", "AJ", "P", "P", "P", None),
    (20, "Fernanda Bulcão Palmeira", None, "A", "A", "A", "A", "A", "A", "A", "A", None),
    (21, "Fernando Gassmann Figueiredo", None, "A", "P", "P", "A", "A", "A", "A", "A", None),
    (22, "Hugo da Cruz Dórea", "*", "P", "P", "P", "P", "P", "P", "P", "P", None),
    (23, "Ítalo Almeida de Moura", None, "P", "P", "P", "P", "P", "P", "P", "AJ", None),
    (24, "Jeidson Antônio Morais Marques", None, "A", "P", "P", "P", "P", "A", "A", "P", None),
    (25, "João Baptista Ferreira", "*", "A", "A", "A", "A", "A", "A", "A", "A", None),
    (26, "João Barreto", "*", "A", "P", "P", "P", "P", "P", "P", "P", None),
    (27, "Jolival Alves Soares", None, "A", "P", "A", "P", "P", "P", "P", "P", None),
    (28, "José Adson Santos Rubem", None, "A", "A", "A", "A", "A", "A", "A", "A", None),
    (29, "José Carlos Oliveira Bispo", None, "A", "P", "A", "A", "A", "A", "A", "A", None),
    (30, "José Carlos Rodrigues", None, "P", "P", "P", "P", "P", "P", "P", "AJ", None),
    (31, "José de Anchieta Leite", "*", "P", "P", "P", "P", "P", "P", "P", "P", None),
    (32, "José Geraldo Lopes Siqueira", None, "P", "P", "P", "P", "P", "P", "A", "P", None),
    (33, "José Raimundo Pereira de Azevedo", "*", "P", "P", "P", "P", "P", "P", "P", "P", None),
    (34, "José Rosa Figueiredo Filho", "*", "A", "P", "A", "A", "P", "P", "P", "A", None),
    (35, "Luiz da Costa Neto", "*", "A", "A", "A", "P", "AJ", "AJ", "P", "A", None),
    (36, "Márcio Lara Medrado", None, "P", "P", "A", "P", "P", "A", "A", "P", None),
    (37, "Maria Clécia Vasconcelos", None, "A", "P", "A", "P", "AJ", "A", "A", "P", "Grafada 'V. de M.F. Costa' na 1a lista - confirmado que e a mesma pessoa"),
    (38, "Miguel Fernandes Dórea", "*", "A", "A", "A", "P", "P", "P", "A", "P", None),
    (39, "Moisés de Santana Silva", None, "P", "P", "A", "P", "P", "P", "P", "AJ", None),
    (40, "Nilson Coelho Lopes", "*", "A", "P", "A", "P", "A", "A", "A", "P", None),
    (41, "Paulo Barreto dos Santos", "*", "P", "P", "P", "AJ", "A", "A", "A", "A", None),
    (43, "Renato Ribeiro da Silva", None, "P", "P", "P", "P", "P", "P", "P", "A", None),
    (44, "Ricardo Gassmann Figueiredo", None, "A", "A", "A", "A", "A", "A", "A", "A", None),
    (46, "Ronald de Freitas Paixão", None, "A", "A", "A", "A", "A", "A", "A", "A", None),
    (47, "Roseli Rodrigues da Silva", None, "P", "P", "P", "P", "P", "AJ", "P", "A", "Grafada 'de Jesus - Rose' na 1a lista"),
    (48, "Ruy Sandes Leal", "*", "A", "P", "P", "P", "P", "P", "A", "P", None),
    (49, "Ruy Sandes Leal Junior", None, "P", "AJ", "AJ", "P", "A", "P", "P", "P", None),
    (50, "Waldécio dos Santos Vita", None, "P", "A", "A", "A", "P", "P", "A", "A", None),
    (51, "Vinícius Guedes Rios", None, "A", "P", "A", "P", "P", "P", "A", "P", None),
    (52, "Yolanda Nati Gassmann Figueiredo", None, "A", "P", "A", "A", "A", "AJ", "A", "A", None),
]

assert all(len(r) == 12 for r in rows), "cada linha precisa ter 12 campos (Nº,Nome,Disp,8 datas,Obs)"
assert len(rows) == 50, f"esperado 50 socios, veio {len(rows)}"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Frequência"

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4A6FA5")
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = ["Nº", "Nome", "Dispensado (*)"] + DATE_LABELS + [
    "Pontos", "Faltas Não Justificadas", "Reuniões Válidas", "% Presença", "Observações",
]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.freeze_panes = "D2"

first_date_col = 4  # D
last_date_col = first_date_col + N_DATES - 1  # K (8 datas)
pontos_col = last_date_col + 1       # L
faltas_col = pontos_col + 1          # M
validas_col = faltas_col + 1         # N
pct_col = validas_col + 1            # O
obs_col = pct_col + 1                # P

for i, r in enumerate(rows):
    excel_row = i + 2
    num, nome, disp = r[0], r[1], r[2]
    dates = r[3:3 + N_DATES]
    obs = r[-1]

    ws.cell(row=excel_row, column=1, value=num)
    ws.cell(row=excel_row, column=2, value=nome)
    ws.cell(row=excel_row, column=3, value=disp)
    for j, code in enumerate(dates):
        ws.cell(row=excel_row, column=first_date_col + j, value=code)

    d_range = f"{get_column_letter(first_date_col)}{excel_row}:{get_column_letter(last_date_col)}{excel_row}"
    ws.cell(row=excel_row, column=pontos_col,
             value=f'=COUNTIF({d_range},"P")+COUNTIF({d_range},"AJ")')
    ws.cell(row=excel_row, column=faltas_col, value=f'=COUNTIF({d_range},"A")')
    ws.cell(row=excel_row, column=validas_col, value=f'={N_DATES}-COUNTIF({d_range},"-")')
    validas_ref = f"{get_column_letter(validas_col)}{excel_row}"
    pontos_ref = f"{get_column_letter(pontos_col)}{excel_row}"
    ws.cell(row=excel_row, column=pct_col,
             value=f'=IF({validas_ref}=0,"-",{pontos_ref}/{validas_ref})')
    ws.cell(row=excel_row, column=obs_col, value=obs)

    for c in range(1, obs_col + 1):
        cell = ws.cell(row=excel_row, column=c)
        cell.border = border
        if c in (first_date_col + j for j in range(N_DATES)):
            cell.alignment = Alignment(horizontal="center")
    ws.cell(row=excel_row, column=pct_col).number_format = "0%"

# larguras de coluna
widths = {1: 5, 2: 34, 3: 6}
for j in range(N_DATES):
    widths[first_date_col + j] = 10
widths[pontos_col] = 9
widths[faltas_col] = 12
widths[validas_col] = 10
widths[pct_col] = 11
widths[obs_col] = 40
for col, w in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# --- aba Legenda ---
leg = wb.create_sheet("Legenda")
leg_rows = [
    ("Código", "Significado"),
    ("P", "Presente (assinou a lista de frequência) — soma 1 ponto"),
    ("AJ", "Ausente com falta justificada (marcou 'Sim' no campo Falta Justificada) — soma 1 ponto, conta como presença"),
    ("A", "Ausente sem justificativa — 0 ponto, conta como falta"),
    ("-", "Reunião em que a pessoa não constava na lista de frequência (ex.: entrou/saiu do quadro)"),
    ("", ""),
    ("(*) na coluna Dispensado", "Sócio dispensado de frequência (Idade + Tempo de Rotary ≥ 85 anos) — pontua normalmente, marcação só informativa"),
    ("", ""),
    ("Decisões e histórico", ""),
    ("07/07/2026", "Primeira lista veio sem a coluna Falta Justificada preenchida em 5 linhas (irregularidade pontual)."),
    ("17/08/2026", "Eliana Mattos de Amorim Bueno: 14/07 e 21/07 corrigidas de A para AJ (falta justificada, doença)."),
    ("18/08/2026", "Ruy Sandes Leal Junior: 28/07 corrigido de A para P (confirmado pelo Jefferson)."),
    ("18/09/2026", "Colunas de agosto (04/08, 11/08, 18/08, 25/08) adicionadas. Leitura feita a partir da lista "
                    "\"Checada\" (com marca-texto amarelo indicando quem esteve presente, marcado por quem coletou "
                    "as assinaturas) — não mais só pela assinatura. Quando o campo \"Falta Justificada: Sim\" está "
                    "marcado, ele é sempre soberano sobre o marca-texto/assinatura (define AJ mesmo se a pessoa "
                    "estiver pintada como presente). A lista pública do site (sem marca-texto) é um documento à parte."),
]
for row in leg_rows:
    leg.append(row)
leg.column_dimensions["A"].width = 26
leg.column_dimensions["B"].width = 110
for c in range(1, 3):
    cell = leg.cell(row=1, column=c)
    cell.font = header_font
    cell.fill = header_fill
for r in range(1, len(leg_rows) + 1):
    for c in (1, 2):
        leg.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

wb.save("Rotary_Frequencia_Mestre.xlsx")
print("OK - planilha gerada com", len(rows), "socios e", N_DATES, "datas.")
